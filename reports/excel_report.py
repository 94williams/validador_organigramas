"""
Genera el Excel de resultados: hoja de detalle (filtrable, con color por
tipo de inconsistencia) + hoja de resumen numérico.

El reporte se adapta al modo de análisis (ver pipeline.ModoAnalisis):
- COMPLETO: incluye columnas de Excel, Word y Organigrama.
- EXCEL_ORGANIGRAMA: excluye por completo las columnas de Word (no solo
  las deja vacías) — Word no participó en el análisis en este modo.

Las columnas de ubicación técnica (página, fila, columna, tabla, celda)
NUNCA se incluyen en el reporte final, sin importar el modo — esa
información se conserva internamente en cada PuestoRecord.ubicacion para
depuración (logs, `Ubicacion.resumen()`), pero no es parte del entregable.
"""
import os
import sys
from collections import Counter
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from models.models import ComparisonResult, TipoInconsistencia  # noqa: E402
from utils.logger import get_logger  # noqa: E402

logger = get_logger("excel_report")

COLOR_POR_TIPO = {
    TipoInconsistencia.OK: "C6EFCE",                          # verde
    TipoInconsistencia.COINCIDE_NORMALIZADO: "D9E1F2",        # azul claro
    TipoInconsistencia.COINCIDE_EQUIVALENCIA: "B4C7E7",       # azul medio (equivalencia institucional)
    TipoInconsistencia.DIFERENCIA_FORMATO: "D9E1F2",
    TipoInconsistencia.POSIBLE_COINCIDENCIA: "FFEB9C",        # amarillo
    TipoInconsistencia.REQUIERE_REVISION: "FFEB9C",
    TipoInconsistencia.COINCIDE_IA: "C9C1F5",                 # lila (resuelto por IA)
    TipoInconsistencia.NIVEL_INCONSISTENTE: "FFC7CE",         # rojo claro
    TipoInconsistencia.NIVEL_FALTANTE: "FFD966",              # naranja
    TipoInconsistencia.NIVEL_FORMATO_INVALIDO: "FFD966",
    TipoInconsistencia.PUESTO_FALTANTE: "FFC7CE",
    TipoInconsistencia.PUESTO_ADICIONAL: "F4B183",            # naranja fuerte
    TipoInconsistencia.DUPLICADO: "E2C6F0",                   # morado claro
    TipoInconsistencia.DUPLICADO_NIVEL_DISTINTO: "C586DB",
    TipoInconsistencia.ERROR_EXTRACCION: "808080",            # gris
}

_ENCABEZADOS_INICIO = [
    "Puesto (clave normalizada)", "Puesto canónico",
    "Tipo de puesto", "Nombre específico", "Niveles según catálogo",
    "Excel - Original", "Excel - Nivel", "Excel - Estado de nivel",
]
_ENCABEZADOS_WORD = [
    "Word - Original", "Word - Nivel", "Word - Estado de nivel",
]
_ENCABEZADOS_FIN = [
    "Organigrama - Original", "Organigrama - Nivel", "Organigrama - Estado de nivel",
    "Resultado", "Tipo de inconsistencia", "Nivel de coincidencia de nombre", "Método de coincidencia",
    "Confianza (%)", "Motivo", "Equivalencia institucional aplicada",
    "Advertencia de catálogo (informativa)",
]


def _incluye_word(modo) -> bool:
    """modo puede ser pipeline.ModoAnalisis o un string equivalente ('completo')."""
    valor = modo.value if hasattr(modo, "value") else str(modo)
    return valor != "excel_organigrama"


def _encabezados(modo) -> List[str]:
    if _incluye_word(modo):
        return _ENCABEZADOS_INICIO + _ENCABEZADOS_WORD + _ENCABEZADOS_FIN
    return _ENCABEZADOS_INICIO + _ENCABEZADOS_FIN


def _fila_desde_resultado(r: ComparisonResult, modo) -> list:
    def campo(record, atributo, default=""):
        if record is None:
            return "—"
        valor = getattr(record, atributo)
        return valor if valor not in (None, "") else default

    resultado_general = "OK" if r.tipo_inconsistencia == TipoInconsistencia.OK else "Revisar"

    representativo = r.excel or r.word or r.organigrama
    tipo_puesto = representativo.tipo_puesto_display if representativo else ""
    nombre_especifico = representativo.nombre_especifico if representativo else ""
    niveles_catalogo = ", ".join(str(n) for n in representativo.niveles_catalogo) if representativo and representativo.niveles_catalogo else "—"

    def estado_nivel(record):
        return record.estado_nivel.value if (record is not None and record.estado_nivel is not None) else "—"

    fila = [
        r.puesto_clave_normalizada,
        r.puesto_canonico or r.puesto_clave_normalizada,
        tipo_puesto, nombre_especifico, niveles_catalogo,
        campo(r.excel, "puesto_original"), campo(r.excel, "nivel_original", "N/D"), estado_nivel(r.excel),
    ]
    if _incluye_word(modo):
        fila += [campo(r.word, "puesto_original"), campo(r.word, "nivel_original", "N/D"), estado_nivel(r.word)]
    fila += [
        campo(r.organigrama, "puesto_original"), campo(r.organigrama, "nivel_original", "N/D"), estado_nivel(r.organigrama),
        resultado_general,
        r.tipo_inconsistencia.value,
        r.nivel_comparacion_nombre.value,
        r.metodo_coincidencia.value if r.metodo_coincidencia else "—",
        round(r.confianza_match, 1),
        r.detalle,
        r.equivalencia_aplicada,
        r.advertencia_catalogo,
    ]
    return fila


def _autoajustar_columnas(ws):
    for col_cells in ws.columns:
        longitud_max = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(longitud_max + 2, 12), 60)


def _escribir_hoja_detalle(wb: Workbook, resultados: List[ComparisonResult], modo):
    ws = wb.active
    ws.title = "Detalle"

    encabezados = _encabezados(modo)
    ws.append(encabezados)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="44546A", end_color="44546A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in resultados:
        fila = _fila_desde_resultado(r, modo)
        ws.append(fila)
        fila_excel_idx = ws.max_row
        color = COLOR_POR_TIPO.get(r.tipo_inconsistencia, "FFFFFF")
        for col_idx in range(1, len(encabezados) + 1):
            ws.cell(row=fila_excel_idx, column=col_idx).fill = PatternFill(
                start_color=color, end_color=color, fill_type="solid",
            )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{ws.max_row}"
    _autoajustar_columnas(ws)


def _escribir_hoja_resumen(wb: Workbook, resultados: List[ComparisonResult], modo,
                            excel_records=None, word_records=None, organigrama_records=None):
    ws = wb.create_sheet("Resumen")
    incluye_word = _incluye_word(modo)

    principales = [
        r for r in resultados
        if not r.puesto_clave_normalizada.startswith("[DUP-") and r.puesto_clave_normalizada != "[ERROR]"
    ]

    total_excel = sum(1 for r in principales if r.excel is not None)
    total_org = sum(1 for r in principales if r.organigrama is not None)
    faltantes_org = sum(1 for r in principales if r.excel and not r.organigrama and r.tipo_inconsistencia == TipoInconsistencia.PUESTO_FALTANTE)
    posibles_coincidencias = sum(1 for r in principales if r.tipo_inconsistencia == TipoInconsistencia.POSIBLE_COINCIDENCIA)
    equivalencias_institucionales = sum(1 for r in principales if r.tipo_inconsistencia == TipoInconsistencia.COINCIDE_EQUIVALENCIA)
    coincidencias_ia = sum(1 for r in principales if r.tipo_inconsistencia == TipoInconsistencia.COINCIDE_IA)
    revision_manual = sum(1 for r in principales if r.tipo_inconsistencia == TipoInconsistencia.REQUIERE_REVISION)
    adicionales_org = sum(1 for r in principales if r.organigrama and not r.excel)
    duplicados = sum(1 for r in resultados if r.puesto_clave_normalizada.startswith("[DUP-"))
    errores = sum(1 for r in resultados if r.puesto_clave_normalizada == "[ERROR]")

    conteo_tipos = Counter(r.tipo_inconsistencia.value for r in resultados)

    ws.append(["Resumen general del análisis"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Modo de análisis: {'Excel + Word + Organigrama (completo)' if incluye_word else 'Excel vs Organigrama (Word no participó)'}"])
    ws.append([])
    ws.append(["Métrica", "Valor"])
    for cell in ws[4]:
        cell.font = Font(bold=True)

    filas_resumen = [("Total de puestos esperados (Excel)", total_excel)]
    if incluye_word:
        total_word = sum(1 for r in principales if r.word is not None)
        en_las_tres = sum(1 for r in principales if r.excel and r.word and r.organigrama)
        faltantes_word = sum(1 for r in principales if r.excel and not r.word and r.tipo_inconsistencia == TipoInconsistencia.PUESTO_FALTANTE)
        adicionales_word = sum(1 for r in principales if r.word and not r.excel)
        filas_resumen += [
            ("Encontrados en Word", total_word),
            ("Encontrados en Organigrama", total_org),
            ("Coincidencias en las tres fuentes", en_las_tres),
        ]
    else:
        filas_resumen.append(("Encontrados en Organigrama", total_org))

    filas_resumen.append(("Coincidencias mediante equivalencia institucional (abreviaturas)", equivalencias_institucionales))
    if coincidencias_ia:
        filas_resumen.append(("Coincidencias resueltas por IA (Ollama)", coincidencias_ia))
    if incluye_word:
        filas_resumen.append(("Faltantes reales en Word", faltantes_word))
    filas_resumen += [
        ("Faltantes reales en Organigrama", faltantes_org),
        ("Posibles coincidencias (revisar)", posibles_coincidencias),
        ("Requieren revisión manual (ambiguos)", revision_manual),
    ]
    if incluye_word:
        filas_resumen.append(("Puestos adicionales en Word (no están en Excel)", adicionales_word))
    filas_resumen += [
        ("Puestos adicionales en Organigrama (no están en Excel)", adicionales_org),
        ("Duplicados detectados", duplicados),
        ("Errores de extracción", errores),
    ]
    for etiqueta, valor in filas_resumen:
        ws.append([etiqueta, valor])

    ws.append([])
    ws.append(["Tipo de inconsistencia", "Cantidad"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for tipo, cantidad in sorted(conteo_tipos.items(), key=lambda x: -x[1]):
        ws.append([tipo, cantidad])

    ws.append([])
    ws.append(["Desglose por fuente"])
    ws[f"A{ws.max_row}"].font = Font(bold=True, size=12)
    ws.append(["Fuente", "Extraídos", "Válidos", "Coincidencias", "Posibles coincidencias", "Faltantes"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    def _fila_fuente(nombre, records, campo_fuente):
        extraidos = len(records) if records else 0
        validos = sum(1 for r in records if r.valido) if records else 0
        coincidencias = sum(
            1 for r in principales
            if getattr(r, campo_fuente) is not None
            and r.tipo_inconsistencia in (TipoInconsistencia.OK, TipoInconsistencia.COINCIDE_NORMALIZADO,
                                           TipoInconsistencia.COINCIDE_EQUIVALENCIA, TipoInconsistencia.COINCIDE_IA)
        )
        posibles = sum(
            1 for r in principales
            if getattr(r, campo_fuente) is not None
            and r.tipo_inconsistencia in (TipoInconsistencia.POSIBLE_COINCIDENCIA, TipoInconsistencia.REQUIERE_REVISION)
        )
        faltantes = sum(
            1 for r in principales
            if getattr(r, campo_fuente) is None and r.excel is not None
            and r.tipo_inconsistencia == TipoInconsistencia.PUESTO_FALTANTE
        )
        ws.append([nombre, extraidos, validos, coincidencias, posibles, faltantes])

    if excel_records is not None:
        _fila_fuente("Excel", excel_records, "excel")
    if incluye_word and word_records is not None:
        _fila_fuente("Word", word_records, "word")
    if organigrama_records is not None:
        _fila_fuente("Organigrama", organigrama_records, "organigrama")

    _autoajustar_columnas(ws)


def generar_reporte(resultados: List[ComparisonResult], ruta_salida: str,
                     excel_records=None, word_records=None, organigrama_records=None,
                     modo: Optional[object] = None) -> str:
    modo = modo or "completo"
    wb = Workbook()
    _escribir_hoja_detalle(wb, resultados, modo)
    _escribir_hoja_resumen(wb, resultados, modo, excel_records, word_records, organigrama_records)
    wb.save(ruta_salida)
    logger.info(f"Reporte generado en: {ruta_salida}")
    return ruta_salida
