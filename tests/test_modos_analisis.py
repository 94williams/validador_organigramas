"""
Pruebas del modo de análisis (pipeline.ModoAnalisis): confirma que
EXCEL_ORGANIGRAMA excluye a Word por completo (nunca lo requiere, nunca lo
compara, nunca lo reporta como faltante), y que COMPLETO sigue funcionando
exactamente igual que antes.
"""
import os
import sys

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from models.models import PuestoRecord, Fuente, Ubicacion, TipoInconsistencia  # noqa: E402
from normalization.normalizer import normalizar_puesto, normalizar_nivel  # noqa: E402
from comparison.comparator import comparar_fuentes  # noqa: E402
from comparison.catalogo_niveles import enriquecer_con_catalogo  # noqa: E402
from pipeline import (  # noqa: E402
    ejecutar_analisis, exportar_reporte, ModoAnalisis,
    _filtrar_rotulos_administrativos,
)
from reports.excel_report import _encabezados  # noqa: E402

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")
RUTA_EXCEL = os.path.join(FIXTURES, "referencia.xlsx")
RUTA_ORG = os.path.join(FIXTURES, "organigrama.pdf")
RUTA_WORD = os.path.join(FIXTURES, "documento_word.docx")


def _rec(fuente, puesto, nivel):
    r = PuestoRecord(
        fuente=fuente, puesto_original=puesto, nivel_original=nivel,
        puesto_normalizado=normalizar_puesto(puesto), nivel_normalizado=normalizar_nivel(nivel),
        ubicacion=Ubicacion(archivo="test"),
    )
    enriquecer_con_catalogo(r)
    return r


def test_filtro_administrativo_aplica_a_todas_las_fuentes():
    textos_ruido = [
        "TOTAL",
        "TOTAL ACTUAL 10",
        "INTEGRÓ: Juan Pérez",
        "TITULAR DEL ÁREA DE ADMINISTRACIÓN: María López",
        "TITULAR DE UNIDAD ADMINISTRATIVA",
        "VARIACIÓN EN COSTO $123.00",
        "NO. PLAZAS ACTUAL 35",
        "NO. PLAZAS PROPUESTA 40",
        "VARIACIÓN EN PLAZAS 5",
    ]
    for fuente in (Fuente.EXCEL, Fuente.WORD, Fuente.ORGANIGRAMA):
        records = [_rec(fuente, texto, "25") for texto in textos_ruido]
        assert _filtrar_rotulos_administrativos(records) == []


def test_filtro_administrativo_conserva_puestos_reales():
    reales = [
        _rec(Fuente.EXCEL, "Dirección de Administración", "40"),
        _rec(Fuente.WORD, "Jefatura de Unidad Departamental de Recursos Humanos", "25"),
        _rec(Fuente.ORGANIGRAMA, "L.C.P. de Planeación", "24"),
    ]
    filtrados = _filtrar_rotulos_administrativos(reales)
    assert [r.puesto_original for r in filtrados] == [r.puesto_original for r in reales]


# ---------------------------------------------------------------------------
# Bug diagnosticado: sin incluir_word=False, un puesto que SÍ coincide entre
# Excel y Organigrama se marcaba falsamente "faltante" solo porque Word
# (que ni siquiera participaba) aparecía como "no encontrado".
# ---------------------------------------------------------------------------
def test_modo_excel_organigrama_no_reporta_falta_en_word():
    excel = [_rec(Fuente.EXCEL, "Jefe de Recursos Humanos", "25")]
    org = [_rec(Fuente.ORGANIGRAMA, "Jefe de Recursos Humanos", "25")]

    resultados = comparar_fuentes(excel, [], org, incluir_word=False)
    r = resultados[0]
    assert r.tipo_inconsistencia == TipoInconsistencia.OK
    assert "Word" not in r.detalle


def test_modo_completo_sigue_reportando_falta_en_word_normalmente():
    excel = [_rec(Fuente.EXCEL, "Jefe de Recursos Humanos", "25")]
    org = [_rec(Fuente.ORGANIGRAMA, "Jefe de Recursos Humanos", "25")]

    resultados = comparar_fuentes(excel, [], org, incluir_word=True)
    r = resultados[0]
    assert r.tipo_inconsistencia == TipoInconsistencia.PUESTO_FALTANTE
    assert "Word" in r.detalle


def test_pipeline_modo_excel_organigrama_no_requiere_word():
    resultado = ejecutar_analisis(
        ruta_excel=RUTA_EXCEL, ruta_organigrama_pdf=RUTA_ORG, ruta_word=None,
        modo=ModoAnalisis.EXCEL_ORGANIGRAMA,
    )
    assert resultado["errores_fatales"] == []
    assert resultado["word_records"] == []
    assert resultado["modo"] == ModoAnalisis.EXCEL_ORGANIGRAMA
    assert all(r.word is None for r in resultado["resultados"])


def test_pipeline_modo_completo_sigue_igual():
    resultado = ejecutar_analisis(
        ruta_excel=RUTA_EXCEL, ruta_organigrama_pdf=RUTA_ORG, ruta_word=RUTA_WORD,
        modo=ModoAnalisis.COMPLETO,
    )
    assert resultado["errores_fatales"] == []
    assert len(resultado["word_records"]) > 0
    assert any(r.word is not None for r in resultado["resultados"])


def test_reporte_modo_excel_organigrama_no_incluye_columnas_de_word(tmp_path):
    resultado = ejecutar_analisis(
        ruta_excel=RUTA_EXCEL, ruta_organigrama_pdf=RUTA_ORG, ruta_word=None,
        modo=ModoAnalisis.EXCEL_ORGANIGRAMA,
    )
    ruta_salida = str(tmp_path / "reporte.xlsx")
    exportar_reporte(
        resultado["resultados"], ruta_salida,
        resultado["excel_records"], resultado["word_records"], resultado["organigrama_records"],
        modo=ModoAnalisis.EXCEL_ORGANIGRAMA,
    )
    import openpyxl
    wb = openpyxl.load_workbook(ruta_salida)
    encabezados = [c.value for c in wb["Detalle"][1]]
    assert not any("Word" in str(e) for e in encabezados)
    assert not any(str(wb["Resumen"].cell(row=i, column=1).value or "").strip() == "Encontrados en Word"
                   for i in range(1, wb["Resumen"].max_row + 1))


def test_reporte_modo_completo_si_incluye_columnas_de_word(tmp_path):
    resultado = ejecutar_analisis(
        ruta_excel=RUTA_EXCEL, ruta_organigrama_pdf=RUTA_ORG, ruta_word=RUTA_WORD,
        modo=ModoAnalisis.COMPLETO,
    )
    ruta_salida = str(tmp_path / "reporte.xlsx")
    exportar_reporte(
        resultado["resultados"], ruta_salida,
        resultado["excel_records"], resultado["word_records"], resultado["organigrama_records"],
        modo=ModoAnalisis.COMPLETO,
    )
    import openpyxl
    wb = openpyxl.load_workbook(ruta_salida)
    encabezados = [c.value for c in wb["Detalle"][1]]
    assert any("Word" in str(e) for e in encabezados)


def test_reporte_nunca_incluye_columnas_de_ubicacion():
    palabras_prohibidas = ["página", "pagina", "fila", "columna", "celda", "tabla", "coordenada", "ubicación", "ubicacion", "ruta"]
    for modo in (ModoAnalisis.COMPLETO, ModoAnalisis.EXCEL_ORGANIGRAMA):
        for encabezado in _encabezados(modo):
            normalizado = encabezado.lower()
            assert not any(p in normalizado for p in palabras_prohibidas), f"'{encabezado}' parece de ubicación (modo={modo})"


def test_jud_lcp_siguen_normalizando_en_ambos_modos():
    excel = [_rec(Fuente.EXCEL, "Líder Coordinador de Proyectos de Planeación", "23")]
    org = [_rec(Fuente.ORGANIGRAMA, "L.C.P. de Planeación", "24")]
    resultados = comparar_fuentes(excel, [], org, incluir_word=False)
    r = resultados[0]
    assert r.organigrama is not None
    assert r.tipo_inconsistencia == TipoInconsistencia.NIVEL_INCONSISTENTE
    assert "LCP" in r.equivalencia_aplicada
