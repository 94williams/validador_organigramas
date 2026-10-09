from copy import deepcopy
from io import BytesIO

import openpyxl
import pytest
from comparison.comparator import comparar_fuentes
from extractors.excel_extractor import extraer_excel
from models.models import Fuente, PuestoRecord, TipoInconsistencia
from normalization.normalizer import normalizar_puesto
from reports.excel_report import _encabezados, _fila_desde_resultado, generar_reporte_bytes


def rec(fuente,nombre,nivel='N25'):
    return PuestoRecord(fuente,nombre,nivel,normalizar_puesto(nombre),nivel)


@pytest.mark.parametrize('excel,word,org',[
 ('Jefatura de Unidad Departamental de Archivo','J.U.D. de Archivo','JUD de Archivo'),
 ('Jefe de Unidad Departamental de Archivo','J U D de Archivo','J.U.D. de Archivo'),
 ('Líder Coordinador de Proyectos de Gestión','l.c.p. de gestion','L C P de Gestión'),
 ('J. U. D. de Archivo','JUD DE ARCHIVO','Jefatura de Unidad Departamental de Archivo'),
])
def test_tres_fuentes_coinciden_tras_homologar(excel,word,org):
    fuentes=([rec(Fuente.EXCEL,excel)],[rec(Fuente.WORD,word)],[rec(Fuente.ORGANIGRAMA,org)])
    antes=deepcopy(fuentes)
    r=comparar_fuentes(*fuentes)[0]
    assert r.similitud_homologada == 100
    assert r.estado_homologacion=='Coinciden'
    assert set(r.comparaciones_homologadas.values())=={'Coinciden'}
    assert r.tipo_inconsistencia==TipoInconsistencia.COINCIDE_EQUIVALENCIA
    assert fuentes==antes


def test_homologacion_no_oculta_nivel_distinto():
    r=comparar_fuentes([rec(Fuente.EXCEL,'Jefatura de Unidad Departamental de Archivo')],[],
                      [rec(Fuente.ORGANIGRAMA,'J.U.D. de Archivo','N29')],incluir_word=False)[0]
    assert r.similitud_homologada == 100
    assert r.estado_homologacion=='Coinciden'
    assert r.tipo_inconsistencia==TipoInconsistencia.NIVEL_INCONSISTENTE


def test_similitud_no_es_igualdad_homologada():
    r=comparar_fuentes([rec(Fuente.EXCEL,'JUD de Seguimiento y Evaluación de Proyectos')],[],
                      [rec(Fuente.ORGANIGRAMA,'JUD de Seguimiento de Proyectos')],incluir_word=False)[0]
    assert 0 < r.similitud_homologada < 100
    assert r.estado_homologacion=='No coinciden'
    assert r.tipo_inconsistencia==TipoInconsistencia.REQUIERE_REVISION


def test_no_confunde_sufijos_distintos():
    resultados=comparar_fuentes([rec(Fuente.EXCEL,'JUD de Archivo A')],[],
                                [rec(Fuente.ORGANIGRAMA,'JUD de Archivo B')],incluir_word=False)
    assert all(r.estado_homologacion!='Coinciden' for r in resultados)


def test_fuente_ausente_y_pareja_presente():
    r=comparar_fuentes([rec(Fuente.EXCEL,'JUD de Archivo')],[],[rec(Fuente.ORGANIGRAMA,'JUD de Archivo')])[0]
    assert r.similitud_homologada is None
    assert r.estado_homologacion=='No comparable'
    assert r.comparaciones_homologadas['Excel vs Organigrama']=='Coinciden'
    assert r.comparaciones_homologadas['Excel vs Word']=='No comparable'
    assert r.tipo_inconsistencia==TipoInconsistencia.PUESTO_FALTANTE


def test_duplicados_no_se_certifican_como_homologacion_univoca():
    resultados=comparar_fuentes([rec(Fuente.EXCEL,'JUD de Archivo'),rec(Fuente.EXCEL,'Jefe de Unidad Departamental de Archivo')],[],
                                [rec(Fuente.ORGANIGRAMA,'JUD de Archivo')],incluir_word=False)
    assert all(r.estado_homologacion!='Coinciden' for r in resultados)
    assert any(r.estado_homologacion=='Revisar duplicados' for r in resultados)
    assert all(r.similitud_homologada is None for r in resultados)


@pytest.mark.parametrize('modo',['completo','excel_organigrama'])
def test_sin_posiciones_y_con_columnas_por_fuente(modo):
    excel=[rec(Fuente.EXCEL,'JUD de Archivo'),rec(Fuente.EXCEL,'LCP de Gestión')]
    org=[rec(Fuente.ORGANIGRAMA,'Líder Coordinador de Proyectos de Gestión'),rec(Fuente.ORGANIGRAMA,'Jefatura de Unidad Departamental de Archivo')]
    word=[rec(Fuente.WORD,'LCP de Gestión'),rec(Fuente.WORD,'J U D de Archivo')]
    resultados=comparar_fuentes(excel,word,org,incluir_word=modo=='completo')
    assert all(r.estado_homologacion=='Coinciden' for r in resultados)
    wb=openpyxl.load_workbook(BytesIO(generar_reporte_bytes(resultados,excel,word,org,modo)))
    headers=[c.value for c in wb['Detalle'][1]]
    assert not any('posición' in h.lower() or 'orden' in h.lower() for h in headers)
    assert headers[-1]=='Similitud de nombres homologados (%)'
    assert ('Word - Puesto homologado' in headers)==(modo=='completo')
    assert ('Homologado: Word vs Organigrama' in headers)==(modo=='completo')
    assert 'Organigrama - Puesto homologado' in headers
    assert len(headers)==len(_fila_desde_resultado(resultados[0],modo))
    assert not any('posición' in str(c.value).lower() for row in wb['Resumen'] for c in row)


def test_excel_no_depende_de_columna_i_y_excluye_encabezado_repetido(tmp_path):
    wb=openpyxl.Workbook();ws=wb.active
    ws.append(['']*8+['ID','Puesto','Nivel'])
    ws.append(['']*8+[None,'JUD de Archivo',25])
    ws.append(['']*8+['ID','Puesto','Nivel'])
    ws.append(['']*8+['texto','LCP de Gestión',None])
    path=tmp_path/'datos.xlsx';wb.save(path)
    records=extraer_excel(str(path))
    assert [r.puesto_original for r in records]==['JUD de Archivo','LCP de Gestión']
    assert all(r.valido for r in records)
    assert records[1].nivel_original is None


def test_porcentaje_toma_peor_pareja_y_se_exporta_como_numero():
    from comparison.homologation import verificar_homologacion, clave_homologada
    from models.models import ComparisonResult
    from rapidfuzz.fuzz import ratio
    e = rec(Fuente.EXCEL, 'JUD de Archivo')
    w = rec(Fuente.WORD, 'JUD de Archivos')
    o = rec(Fuente.ORGANIGRAMA, 'JUD de Archivo General')
    r = ComparisonResult('archivo', excel=e, word=w, organigrama=o)
    verificar_homologacion([r], [e], [w], [o])
    esperado = min(round(ratio(clave_homologada(a), clave_homologada(b)), 1)
                   for a, b in [(e, w), (e, o), (w, o)])
    assert r.similitud_homologada == esperado
    wb = openpyxl.load_workbook(BytesIO(generar_reporte_bytes([r])))
    ws = wb['Detalle']
    headers = [c.value for c in ws[1]]
    i = headers.index('Excel - Puesto homologado')
    assert headers[i:i+3] == ['Excel - Puesto homologado', 'Word - Puesto homologado', 'Organigrama - Puesto homologado']
    assert 'Advertencia de formato (mayúsculas)' not in headers
    cell = ws.cell(2, headers.index('Similitud de nombres homologados (%)') + 1)
    assert cell.data_type == 'n'
    assert cell.value == esperado
    o.error = 'Texto ilegible'
    verificar_homologacion([r], [e], [w], [o])
    assert r.similitud_homologada is None
    fila = _fila_desde_resultado(r, 'completo')
    assert fila[headers.index('Similitud de nombres homologados (%)')] == 'No comparable'
