from copy import deepcopy
from io import BytesIO

import docx
import fitz
import openpyxl
import pytest

from comparison.comparator import comparar_fuentes
from extractors.excel_extractor import extraer_excel, _leer_id
from extractors.organigrama_extractor import extraer_organigrama
from models.models import PuestoRecord, Fuente, TipoInconsistencia
from normalization.normalizer import normalizar_puesto, normalizar_nivel
from pipeline import ejecutar_analisis
from reports.excel_report import generar_reporte_bytes


NOMBRES = ['Dirección de Administración y Finanzas en la Secretaría del Deporte',
           'Subdirección de Administración, Desarrollo de Personal y Recursos Financieros',
           'Líder Coordinador de Proyectos de Desarrollo de Personal',
           'Líder Coordinador de Proyectos de Finanzas',
           'Subdirección de Recursos Materiales y Servicios Generales',
           'Líder Coordinador de Proyectos de Abastecimiento y Servicios']
NIVELES = [40, 29, 24, 24, 29, 24]
PADRES = [None, 0, 1, 1, 0, 4]
CAJAS = [(220, 20, 400, 90), (100, 150, 280, 225), (0, 300, 180, 370),
         (200, 300, 380, 370), (450, 150, 630, 225), (450, 300, 630, 370)]


def crear_pdf(path, nombres=NOMBRES, conectado=True, paginas=1):
    doc = fitz.open()
    for _ in range(paginas):
        page = doc.new_page(width=650, height=410)
        for i in [5, 4, 3, 2, 1, 0]:  # orden interno del PDF deliberadamente distinto
            x0, y0, x1, y1 = CAJAS[i]
            page.draw_rect(fitz.Rect(*CAJAS[i]))
            texto = nombres[i].replace('Líder Coordinador de Proyectos', 'L.C.P.')
            page.insert_textbox(fitz.Rect(x0+3, y0+3, x1-3, y1-3), texto + '\n' + str(NIVELES[i]), fontsize=9)
        if conectado:
            for padre, hijos, altura in [(0, [1, 4], 120), (1, [2, 3], 260), (4, [5], 260)]:
                box = CAJAS[padre]; x = (box[0]+box[2])/2
                xs = [(CAJAS[h][0]+CAJAS[h][2])/2 for h in hijos]
                page.draw_line((x, box[3]), (x, altura))
                if min(xs+[x]) != max(xs+[x]):
                    page.draw_line((min(xs+[x]), altura), (max(xs+[x]), altura))
                for h, hx in zip(hijos, xs):
                    page.draw_line((hx, altura), (hx, CAJAS[h][1]))
    doc.save(path); doc.close()


def crear_excel(path, nombres=NOMBRES, ids=None):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(['']*8 + ['ID', 'Puesto', 'Nivel', 'Puesto del jefe inmediato'])
    for i, nombre in enumerate(nombres):
        if i == 3:
            ws.append(['']*8 + ['ID', 'Puesto', 'Nivel', 'Puesto del jefe inmediato'])
        ws.append(['']*8 + [(ids or list(range(1, 7)))[i], nombre, NIVELES[i],
                          nombres[PADRES[i]] if PADRES[i] is not None else None])
    wb.save(path)


def crear_word(path, nombres=NOMBRES):
    doc = docx.Document(); table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Puesto'; table.cell(0, 1).text = 'Nivel'
    for nombre, nivel in zip(nombres, NIVELES):
        row = table.add_row(); row.cells[0].text = nombre; row.cells[1].text = str(nivel)
    doc.save(path)


def ejecutar(tmp_path, nombres=NOMBRES, word=None, ids=None, conectado=True, paginas=1):
    e, w, o = (tmp_path/n for n in ('base.xlsx', 'base.docx', 'base.pdf'))
    crear_excel(e, nombres, ids); crear_word(w, word or nombres); crear_pdf(o, nombres, conectado, paginas)
    return ejecutar_analisis(str(e), str(o), str(w))


def test_recorrido_ejemplo_y_reporte_sin_advertencia_mayusculas(tmp_path):
    data = ejecutar(tmp_path)
    assert not data['errores_fatales']
    rs = data['resultados']
    assert len(rs) == 6
    assert [r.id_excel for r in data['excel_records']] == list(range(1, 7))
    assert [r.posicion_lectura for r in data['organigrama_records']] == list(range(1, 7))
    assert all(r.estado_orden == 'Misma posición' for r in rs)
    assert all(r.estado_homologacion == 'Coinciden' for r in rs)
    assert rs[3].organigrama.jefe_inmediato == NOMBRES[1]
    assert rs[5].organigrama.jefe_inmediato == NOMBRES[4]
    assert all(r.estado_jefe == 'Coincide' for r in rs[1:])
    wb = openpyxl.load_workbook(BytesIO(generar_reporte_bytes(rs)))
    headers = [c.value for c in wb['Detalle'][1]]
    assert 'Advertencia de formato (mayúsculas)' not in headers
    assert [row[headers.index('ID esperado - Excel')] for row in list(wb['Detalle'].values)[1:]] == list(range(1, 7))
    assert len(headers) == len(next(wb['Detalle'].iter_rows(min_row=2)))


def test_nombre_repetido_se_conserva_por_posicion_y_jefe(tmp_path):
    nombres = NOMBRES.copy(); nombres[2] = nombres[5] = 'Líder Coordinador de Proyectos de Apoyo'
    data = ejecutar(tmp_path, nombres)
    rs = data['resultados']
    assert len(rs) == 6
    assert all(r.estado_orden == 'Misma posición' for r in rs)
    assert all(r.estado_homologacion == 'Coinciden' for r in rs)
    assert 'Word no permite verificar' in rs[2].detalle_orden
    assert rs[2].organigrama is not rs[5].organigrama
    assert rs[2].organigrama.jefe_inmediato != rs[5].organigrama.jefe_inmediato


def test_word_desplazado_muestra_lo_real_y_localiza_esperado(tmp_path):
    word = NOMBRES.copy(); word[2], word[3] = word[3], word[2]
    rs = ejecutar(tmp_path, word=word)['resultados']
    assert rs[2].word.puesto_original == NOMBRES[3]
    assert rs[2].estado_homologacion == 'No coinciden'
    assert rs[2].orden_word == 'Diferente posición'
    assert 'posición 4; esperada 3' in rs[2].detalle_orden
    assert rs[2].orden_organigrama == 'Misma posición'


def test_organigrama_desplazado_usa_el_jefe_para_localizar_repetidos(tmp_path):
    nombres = NOMBRES.copy(); nombres[2] = nombres[5] = 'Líder Coordinador de Proyectos de Apoyo'
    data = ejecutar(tmp_path, nombres)
    excel = data['excel_records']
    # Cambiar los ID de ambos grupos completos, manteniendo cada jefe declarado.
    for r, n in zip(excel, [1, 4, 5, 6, 2, 3]):
        r.id_excel = n
    rs = comparar_fuentes(excel, data['word_records'], data['organigrama_records'])
    r = next(r for r in rs if r.excel and r.excel.id_excel == 3)
    assert r.estado_jefe == 'Jefe inmediato diferente'
    assert r.orden_organigrama == 'Diferente posición'
    assert 'Organigrama: posición 6; esperada 3' in r.detalle_orden


@pytest.mark.parametrize('valor', [None, '', 'PUESTO', 0, -1, 1.5, True, 'NaN', 'Infinity'])
def test_id_invalido(valor):
    assert _leer_id(valor) is None


def test_ids_invalidos_no_consumen_la_misma_posicion(tmp_path):
    rs = ejecutar(tmp_path, ids=[1, 1, None, 4, 5, 6])['resultados']
    pendientes = [r for r in rs if r.excel and r.estado_orden == 'Revisar ID de Excel']
    assert len(pendientes) == 3
    assert all(r.word is None and r.organigrama is None for r in pendientes)
    from reports.excel_report import _detalle_nombres
    assert all('aparición sin asignar' in _detalle_nombres(r, 'completo') for r in pendientes)
    assert len([r for r in rs if r.excel]) == 6


def test_id_11_no_se_renumera_y_nivel_distinto_no_se_oculta(tmp_path):
    data = ejecutar(tmp_path, ids=list(range(11, 17)))
    r = data['resultados'][0]
    assert r.excel.id_excel == 11
    assert r.word is None
    assert 'posición 1; esperada 11' in r.detalle_orden
    # Con el orden corregido, una diferencia de nivel sigue requiriendo ajuste.
    for i, e in enumerate(data['excel_records'], 1):
        e.id_excel = i
    data['word_records'][2].nivel_normalizado = 'N25'
    data['word_records'][2].nivel_original = '25'
    r = comparar_fuentes(data['excel_records'], data['word_records'], data['organigrama_records'])[2]
    assert r.tipo_inconsistencia == TipoInconsistencia.NIVEL_INCONSISTENTE
    assert r.estado_orden == 'Misma posición'


@pytest.mark.parametrize('opciones', [dict(conectado=False), dict(paginas=2)])
def test_no_inventa_orden_sin_conexiones_o_continuidad(tmp_path, opciones):
    data = ejecutar(tmp_path, **opciones)
    rs = [r for r in data['resultados'] if r.excel]
    assert all(r.orden_organigrama == 'No verificable' for r in rs)
    assert all(r.posicion_organigrama is None for r in rs)


def test_dos_fuentes_no_inventa_word_y_no_muta_entradas(tmp_path):
    data = ejecutar(tmp_path)
    antes = deepcopy(data)
    rs = comparar_fuentes(data['excel_records'], [], data['organigrama_records'], incluir_word=False)
    assert all(r.estado_orden == 'Misma posición' for r in rs)
    assert all('Word' not in r.detalle_orden for r in rs)
    assert data == antes


def test_jefe_como_id_se_resuelve_sin_confundir_columna_puesto(tmp_path):
    p = tmp_path/'ids.xlsx'; crear_excel(p)
    w = openpyxl.load_workbook(p); ws = w.active
    ws.cell(3, 12, 1); w.save(p)
    rs = extraer_excel(str(p))
    assert rs[1].puesto_original == NOMBRES[1]
    assert rs[1].jefe_inmediato == NOMBRES[0]


def test_candidatos_repetidos_fuera_de_posicion_son_ambiguos(tmp_path):
    nombres = NOMBRES.copy(); nombres[2] = nombres[3] = 'Líder Coordinador de Proyectos de Apoyo'
    data = ejecutar(tmp_path, nombres)
    data['excel_records'][5].puesto_original = nombres[2]
    data['excel_records'][5].puesto_normalizado = normalizar_puesto(nombres[2])
    rs = comparar_fuentes(data['excel_records'], data['word_records'], data['organigrama_records'])
    assert rs[5].orden_word == 'Coincidencia ambigua'
    assert 'posiciones 3, 4' in rs[5].detalle_orden
    assert rs[5].word.puesto_original == NOMBRES[5]


def test_jefe_con_nombre_repetido_no_certifica_identidad(tmp_path):
    nombres = NOMBRES.copy(); nombres[1] = nombres[4] = 'Subdirección de Administración'
    data = ejecutar(tmp_path, nombres)
    assert data['resultados'][2].estado_jefe == 'Jefe ambiguo'
    assert data['resultados'][2].estado_orden == 'Jefe ambiguo'
    assert data['resultados'][2].tipo_inconsistencia == TipoInconsistencia.REQUIERE_REVISION


def test_id_jefe_inexistente_queda_pendiente(tmp_path):
    data = ejecutar(tmp_path)
    data['excel_records'][1].jefe_inmediato = '[ID de jefe no resoluble: 99]'
    r = comparar_fuentes(data['excel_records'], data['word_records'], data['organigrama_records'])[1]
    assert r.estado_jefe == 'No verificable'
    assert r.estado_orden == 'No verificable'
    assert 'corregir la referencia' in r.detalle_orden


@pytest.mark.parametrize('tipo', ['curva', 'lateral', 'cruce'])
def test_geometria_ambigua_no_inventa_jefes(tmp_path, tipo):
    p = tmp_path/'normal.pdf'; crear_pdf(p)
    doc = fitz.open(p); page = doc[0]
    if tipo == 'curva':
        page.draw_bezier((10, 390), (20, 385), (30, 385), (40, 390))
    elif tipo == 'lateral':
        page.draw_line((280, 190), (450, 190))
    else:
        page.draw_line((350, 100), (350, 140))
    out = tmp_path/'ambiguo.pdf'; doc.save(out); doc.close()
    rs = extraer_organigrama(str(out))
    assert rs and all(r.error_orden and r.posicion_lectura is None for r in rs)
    assert all(not r.jerarquia_verificada for r in rs)


def test_ui_filtra_diferencias_de_orden(tmp_path, monkeypatch):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    import config
    monkeypatch.setattr(config, 'APP_PASSWORD', None)
    nombres_word = NOMBRES.copy(); nombres_word[2], nombres_word[3] = nombres_word[3], nombres_word[2]
    data = ejecutar(tmp_path, word=nombres_word)
    app = AppTest.from_file(str(Path(__file__).parents[1]/'ui/app.py'))
    app.session_state['version_resultados'] = '2026-10-12-L'
    app.session_state['resultado_analisis'] = data
    app.run(timeout=15)
    assert not app.exception
    next(c for c in app.checkbox if c.label == 'Solo diferencias o pendientes de orden').check().run(timeout=15)
    assert not app.exception
    assert len(app.dataframe[0].value) == 2
    assert app.dataframe[0].value['ID esperado - Excel'].tolist() == [3, 4]
