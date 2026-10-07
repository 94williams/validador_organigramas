from io import BytesIO
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from comparison.comparator import comparar_fuentes
from extractors.excel_extractor import extraer_excel
from models.models import Fuente, PuestoRecord, TipoInconsistencia
from normalization.normalizer import normalizar_puesto
from pipeline import ejecutar_analisis, ModoAnalisis
from reports.excel_report import generar_reporte_bytes
from ui import analysis_service


def rec(fuente, nombre='JUD de Seguimiento y Evaluación de Proyectos', nivel='N25'):
    return PuestoRecord(fuente, nombre, nivel, normalizar_puesto(nombre), nivel)


def result(excel, word, org, modo=ModoAnalisis.COMPLETO):
    return dict(excel_records=excel, word_records=word, organigrama_records=org,
                resultados=comparar_fuentes(excel, word, org, incluir_word=modo == ModoAnalisis.COMPLETO),
                modo=modo, errores_fatales=[])


@pytest.mark.parametrize('ausente', ['Word', 'Organigrama'])
@pytest.mark.parametrize('ambiguo', [False, True])
def test_faltante_permanece_en_detalle_y_resumen(ausente, ambiguo):
    excel = [rec(Fuente.EXCEL)]
    otro = rec(Fuente.WORD if ausente == 'Organigrama' else Fuente.ORGANIGRAMA,
               nombre='JUD de Seguimiento de Proyectos' if ambiguo else excel[0].puesto_original,
               nivel='N29')
    word, org = ([otro], []) if ausente == 'Organigrama' else ([], [otro])
    data = result(excel, word, org)
    hallazgo = data['resultados'][0]
    assert ausente in hallazgo.detalle
    assert hallazgo.tipo_inconsistencia == (TipoInconsistencia.REQUIERE_REVISION if ambiguo else TipoInconsistencia.NIVEL_INCONSISTENTE)
    wb = openpyxl.load_workbook(BytesIO(generar_reporte_bytes(data['resultados'], excel, word, org)))
    filas = list(wb['Resumen'].values)
    assert next(r[1] for r in filas if r[0] == f'Faltantes reales en {ausente}') == 1
    assert next(r[5] for r in filas if r[0] == ausente) == 1
    assert ausente in next(str(c.value) for row in wb['Detalle'] for c in row if c.value == hallazgo.detalle)


def test_hoja_inexistente_marca_pipeline_incompleto(tmp_path, monkeypatch):
    archivo = tmp_path / 'input.xlsx'
    wb = openpyxl.Workbook(); wb.active.title = 'Puestos'
    wb.active.append(['Puesto', 'Nivel']); wb.active.append(['Dirección de Administración', 40]); wb.save(archivo)
    with pytest.raises(ValueError, match='Hojas disponibles: Puestos'):
        extraer_excel(str(archivo), hoja='Incorrecta')
    monkeypatch.setattr('pipeline.extraer_organigrama', lambda _: [rec(Fuente.ORGANIGRAMA)])
    data = ejecutar_analisis(str(archivo), 'unused.pdf', hoja_excel='Incorrecta', modo=ModoAnalisis.EXCEL_ORGANIGRAMA)
    assert data['errores_fatales'] and 'Análisis incompleto' in data['errores_fatales'][0]
    assert any(r.tipo_inconsistencia == TipoInconsistencia.ERROR_EXTRACCION for r in data['resultados'])


def uploaded(nombre):
    archivo = BytesIO(b'contenido simulado'); archivo.name = nombre
    return archivo


def test_dos_usuarios_concurrentes_no_comparten_reporte(tmp_path, monkeypatch):
    barrier = Barrier(2)
    rutas = []
    def pipeline(**kwargs):
        path = Path(kwargs['ruta_excel']); rutas.append(path)
        assert path.exists()
        barrier.wait(timeout=10)
        nombre = kwargs['hoja_excel']
        return result([rec(Fuente.EXCEL, nombre)], [], [rec(Fuente.ORGANIGRAMA, nombre)], ModoAnalisis.EXCEL_ORGANIGRAMA)
    monkeypatch.setattr(analysis_service, 'ejecutar_analisis', pipeline)
    def run(nombre):
        return analysis_service.analizar_archivos(uploaded('a.xlsx'), uploaded('b.pdf'), hoja=nombre, modo=ModoAnalisis.EXCEL_ORGANIGRAMA)[1]
    with ThreadPoolExecutor(max_workers=2) as pool:
        reportes = list(pool.map(run, ['PUESTO USUARIO A', 'PUESTO USUARIO B']))
    assert rutas[0].parent != rutas[1].parent
    assert all(not path.parent.exists() for path in rutas)
    for nombre, contenido in zip(['PUESTO USUARIO A', 'PUESTO USUARIO B'], reportes):
        wb = openpyxl.load_workbook(BytesIO(contenido))
        assert wb['Detalle']['F2'].value == nombre
        assert not any('Word -' in str(c.value) for c in wb['Detalle'][1])


def test_temporales_se_limpian_si_falla_exportacion(monkeypatch):
    rutas = []
    def pipeline(**kwargs):
        rutas.append(Path(kwargs['ruta_excel']))
        return result([], [], [])
    def fail(*args, **kwargs):
        raise OSError('No se puede exportar')
    monkeypatch.setattr(analysis_service, 'ejecutar_analisis', pipeline)
    monkeypatch.setattr(analysis_service, 'generar_reporte_bytes', fail)
    with pytest.raises(OSError):
        analysis_service.analizar_archivos(uploaded('a.xlsx'), uploaded('b.pdf'), uploaded('c.docx'))
    assert not rutas[0].parent.exists()


def test_ambiguedad_permanece_en_revision_manual():
    data = result([rec(Fuente.EXCEL)], [rec(Fuente.WORD)], [rec(Fuente.ORGANIGRAMA, 'JUD de Seguimiento de Proyectos')])
    assert data['resultados'][0].tipo_inconsistencia == TipoInconsistencia.REQUIERE_REVISION


@pytest.mark.parametrize('busqueda', ['(', '[', '*', 'seguimiento'])
def test_ui_busqueda_literal_y_conteos_reales(busqueda, monkeypatch):
    monkeypatch.setattr('config.APP_PASSWORD', None)
    excel = [rec(Fuente.EXCEL), rec(Fuente.EXCEL), PuestoRecord(Fuente.EXCEL, '', None, error='ilegible')]
    data = result(excel, [], [rec(Fuente.ORGANIGRAMA)], ModoAnalisis.EXCEL_ORGANIGRAMA)
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'ui/app.py'))
    app.session_state['resultado_analisis'] = data
    app.session_state['reporte_bytes'] = generar_reporte_bytes(data['resultados'], modo=data['modo'])
    app.run(timeout=15)
    assert not app.exception
    assert next(m.value for m in app.metric if m.label == 'Puestos en Excel') == '2'
    campo = next(t for t in app.text_input if t.label == 'Buscar puesto')
    campo.set_value(busqueda).run(timeout=15)
    assert not app.exception
    assert len(app.dataframe[0].value) == ({'seguimiento': 3, '[': 3}.get(busqueda, 0))
