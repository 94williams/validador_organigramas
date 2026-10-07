from copy import deepcopy
from io import BytesIO

import openpyxl
import pytest
from models.models import ComparisonResult, Fuente, PuestoRecord
from normalization.display_name import homologar_puesto_excel
from reports.excel_report import _encabezados, _fila_desde_resultado, generar_reporte_bytes


@pytest.mark.parametrize('original,esperado', [
    ('Jefatura de Unidad Departamental de Atención Ciudadana', 'J.U.D. de Atención Ciudadana'),
    ('Jefe de Unidad Departamental de Archivo', 'J.U.D. de Archivo'),
    ('Líder Coordinador de Proyectos de Evaluación', 'L.C.P. de Evaluación'),
    ('LIDER COORDINADOR DE PROYECTOS A', 'L.C.P. A'),
    ('jud de Recursos Humanos', 'J.U.D. de Recursos Humanos'),
    ('l.c.p. de Planeación', 'L.C.P. de Planeación'),
    (' J. U. D. de Obras Públicas "A"', 'J.U.D. de Obras Públicas "A"'),
    ('L C P de Apoyo', 'L.C.P. de Apoyo'),
    ('J.U.D.', 'J.U.D.'),
    ('Dirección de Supervisión de J.U.D.', 'Dirección de Supervisión de J.U.D.'),
    ('Judicial', 'Judicial'),
    ('', ''),
])
def test_homologacion_conserva_nombre_especifico(original, esperado):
    assert homologar_puesto_excel(original) == esperado
    assert homologar_puesto_excel(esperado) == esperado


@pytest.mark.parametrize('modo', ['completo', 'excel_organigrama'])
def test_ultima_columna_y_original_sin_modificaciones(modo):
    original = 'Líder Coordinador de Proyectos de Gestión'
    r = ComparisonResult('clave', excel=PuestoRecord(Fuente.EXCEL, original, '24'))
    antes = deepcopy(r)
    headers = _encabezados(modo)
    fila = _fila_desde_resultado(r, modo)
    assert headers[-2:] == ['Observación de orden', 'Excel - Puesto homologado']
    assert len(headers) == len(fila)
    assert fila[headers.index('Excel - Original')] == original
    assert fila[-1] == 'L.C.P. de Gestión'
    wb = openpyxl.load_workbook(BytesIO(generar_reporte_bytes([r], modo=modo)))
    ws = wb['Detalle']
    assert ws.cell(1, ws.max_column).value == 'Excel - Puesto homologado'
    assert ws.cell(2, ws.max_column).value == 'L.C.P. de Gestión'
    assert r == antes


def test_sin_excel_no_inventa_nombre_desde_otra_fuente():
    r = ComparisonResult('clave', word=PuestoRecord(Fuente.WORD, 'JUD de Archivo', '25'))
    assert _fila_desde_resultado(r, 'completo')[-1] == '—'
