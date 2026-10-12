from io import BytesIO

import openpyxl
import pytest

from comparison.comparator import comparar_fuentes
from extractors.excel_extractor import extraer_excel
from reports.excel_report import generar_reporte_bytes


JEFE = 'Dirección de Administración y Finanzas'


def libro(encabezado='Responsable superior', jefe=JEFE):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append([''] * 8 + ['ID', 'Puesto', 'Nivel', encabezado])
    ws.append([''] * 8 + [1, JEFE, 40, None])
    ws.append([''] * 8 + [2, 'Subdirección de Administración', 29, jefe])
    return wb


@pytest.mark.parametrize('encabezado', [None, 'Responsable superior', 'Jefe\ninmediato (puesto)',
                                      'DENOMINACIÓN DEL PUESTO DEL JEFE INMEDIATO'])
def test_jefe_en_l_se_exporta_con_encabezado_variable(tmp_path, encabezado):
    wb = libro(encabezado); p = tmp_path/'entrada.xlsx'; wb.save(p)
    rs = extraer_excel(str(p))
    assert rs[0].jefe_inmediato is None
    assert rs[1].jefe_inmediato == JEFE
    assert rs[1].puesto_original == 'Subdirección de Administración'
    assert rs[1].nivel_normalizado == 'N29'
    resultados = comparar_fuentes(rs, [], [], incluir_word=False)
    report = openpyxl.load_workbook(BytesIO(generar_reporte_bytes(resultados, modo='excel_organigrama')))
    ws = report['Detalle']; headers = [c.value for c in ws[1]]
    assert ws.cell(3, headers.index('Jefe inmediato - Excel') + 1).value == JEFE


def test_encabezado_jefe_otra_fila_y_combinado(tmp_path):
    wb = libro(); ws = wb.active
    ws.insert_rows(1)
    ws['L1'] = 'Jefe inmediato'
    ws.merge_cells('L1:L2')
    p = tmp_path/'entrada.xlsx'; wb.save(p)
    assert extraer_excel(str(p))[1].jefe_inmediato == JEFE


def test_celdas_jefe_combinadas_sin_rellenar_vacios_normales(tmp_path):
    wb = libro(); ws = wb.active
    ws.append([''] * 8 + [3, 'LCP de Finanzas', 24, None])
    ws.append([''] * 8 + [4, 'LCP de Archivo', 24, None])
    ws.merge_cells('L3:L4')
    p = tmp_path/'entrada.xlsx'; wb.save(p)
    rs = extraer_excel(str(p))
    assert [r.jefe_inmediato for r in rs] == [None, JEFE, JEFE, None]


def test_encabezado_en_otra_columna_tiene_prioridad_sobre_l(tmp_path):
    wb = libro(); ws = wb.active
    ws['M1'] = 'Nombre del jefe inmediato'; ws['M3'] = 'Dirección de Operación'
    p = tmp_path/'entrada.xlsx'; wb.save(p)
    assert extraer_excel(str(p))[1].jefe_inmediato == 'Dirección de Operación'


def test_respaldo_l_no_reemplaza_puesto_ni_nivel(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append([''] * 8 + ['ID', 'Nota', 'Nivel', 'Puesto'])
    ws.append([''] * 8 + [1, 'Referencia', 25, 'JUD de Archivo'])
    p = tmp_path/'entrada.xlsx'; wb.save(p)
    r = extraer_excel(str(p))[0]
    assert r.puesto_original == 'JUD de Archivo'
    assert r.jefe_inmediato is None


def test_jefe_en_l_puede_ser_id_sin_encabezado_reconocido(tmp_path):
    wb = libro(jefe=1); p = tmp_path/'entrada.xlsx'; wb.save(p)
    assert extraer_excel(str(p))[1].jefe_inmediato == JEFE
