import os
import sys

import openpyxl

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from extractors.excel_extractor import extraer_excel  # noqa: E402


def test_excel_reconoce_denominacion_del_puesto(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Puestos"
    ws["A1"] = "CANTIDAD"
    ws["B1"] = "DENOMINACION DEL PUESTO"
    ws["C1"] = "NIVEL"
    ws["A2"] = "1"
    ws["B2"] = "Jefe de Departamento"
    ws["C2"] = 10

    ruta = tmp_path / "excel_puestos.xlsx"
    wb.save(ruta)

    registros = extraer_excel(str(ruta))

    assert len(registros) == 1
    assert registros[0].puesto_original == "Jefe de Departamento"
    assert registros[0].nivel_original == "10"


def test_excel_ignora_rotulos_administrativos_y_totales_en_columna_j(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Puestos"
    ws["J1"] = "PUESTO"
    ws["K1"] = "NIVEL"

    valores_ignorados = [
        "INTEGRÓ",
        "Titular de Unidad Administrativa",
        "Titular del Área de Administración",
        "TOTAL",
        "TOTAL ACTUAL",
        "TOTAL PROPUESTO",
        "VARIACIÓN EN COSTO",
        "NO. PLAZAS ACTUAL",
        "NO. PLAZAS PROPUESTA",
        "VARIACIÓN EN PLAZAS",
    ]

    fila = 2
    for valor in valores_ignorados:
        ws.cell(row=fila, column=10, value=valor)
        ws.cell(row=fila, column=11, value="25")
        fila += 1

    ws.cell(row=fila, column=10, value="Jefatura de Unidad Departamental de Recursos Humanos")
    ws.cell(row=fila, column=11, value="25")

    ruta = tmp_path / "excel_con_totales.xlsx"
    wb.save(ruta)

    registros = extraer_excel(str(ruta))

    assert len(registros) == 1
    assert registros[0].puesto_original == "Jefatura de Unidad Departamental de Recursos Humanos"
    assert registros[0].nivel_original == "25"


def test_excel_ignora_variantes_y_texto_adicional_de_rotulos(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Puestos"
    ws["J1"] = "PUESTO"
    ws["K1"] = "NIVEL"

    variantes = [
        "  integró: Juan Pérez  ",
        "TITULAR DE UNIDAD ADMINISTRATIVA - DIRECCIÓN X",
        "Titular del area de administracion: María López",
        "TOTAL 125",
        "TOTAL ACTUAL: 50",
        "variacion en costo $123.00",
        "NO PLAZAS ACTUAL 30",
        "Variación   en   plazas: 2",
    ]

    for idx, valor in enumerate(variantes, start=2):
        ws.cell(row=idx, column=10, value=valor)
        ws.cell(row=idx, column=11, value="20")

    fila_real = len(variantes) + 2
    ws.cell(row=fila_real, column=10, value="Dirección de Administración")
    ws.cell(row=fila_real, column=11, value="40")

    ruta = tmp_path / "excel_variantes_totales.xlsx"
    wb.save(ruta)

    registros = extraer_excel(str(ruta))

    assert len(registros) == 1
    assert registros[0].puesto_original == "Dirección de Administración"
