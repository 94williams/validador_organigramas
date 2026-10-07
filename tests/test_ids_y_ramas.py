from io import BytesIO
from pathlib import Path

import docx
import fitz
import openpyxl
import pytest

from extractors.excel_extractor import extraer_excel, _leer_consecutivo
from extractors.organigrama_extractor import extraer_organigrama
from comparison.comparator import comparar_fuentes
from models.models import Fuente, PuestoRecord, TipoInconsistencia
from pipeline import ejecutar_analisis
from reports.excel_report import generar_reporte_bytes

NOMBRES = ['Dirección General', 'Dirección de Administración', 'JUD de Recursos Humanos',
           'LCP de Capacitación', 'JUD de Finanzas', 'Dirección de Operación',
           'Subdirección de Supervisión', 'JUD de Atención']
# Orden de dibujo intencionalmente distinto al recorrido por ramas.
CAJAS = [(300,20,460,60), (80,140,240,180), (0,260,160,300), (0,380,160,420),
         (170,260,330,300), (470,140,630,180), (390,260,550,300), (560,260,720,300)]


def crear_pdf(path, conectado=True, multipagina=False, curva=False):
    doc = fitz.open();page=doc.new_page(width=750,height=500)
    for i in [7,4,1,6,0,5,3,2]:
        x0,y0,x1,y1=CAJAS[i]
        page.draw_rect(fitz.Rect(*CAJAS[i]))
        page.insert_textbox(fitz.Rect(x0+3,y0+3,x1-3,y1-3),NOMBRES[i]+'\n25',fontsize=8)
    if conectado:
        for padre,hijos,y in [(0,[1,5],100),(1,[2,4],220),(2,[3],350),(5,[6,7],220)]:
            caja=CAJAS[padre];x=(caja[0]+caja[2])/2
            xs=[(CAJAS[h][0]+CAJAS[h][2])/2 for h in hijos]
            page.draw_line((x,caja[3]),(x,y))
            if min([x]+xs)!=max([x]+xs):page.draw_line((min([x]+xs),y),(max([x]+xs),y))
            for h,hx in zip(hijos,xs):page.draw_line((hx,y),(hx,CAJAS[h][1]))
    if curva:page.draw_bezier((10,450),(20,430),(30,430),(40,450))
    if multipagina:
        page=doc.new_page(width=750,height=500)
        page.draw_rect(fitz.Rect(20,20,250,80));page.insert_text((25,45),'Enlace adicional 25')
    doc.save(path);doc.close()


def crear_excel(path, ids=None, repetir_header=False):
    wb=openpyxl.Workbook();ws=wb.active
    ws.append(['']*8+['Consecutivo','Puesto','Nivel'])
    for i in range(8):
        if repetir_header and i==3:ws.append(['']*8+['Consecutivo','Puesto','Nivel'])
        ws.append(['']*8+[(ids or list(range(1,9)))[i],NOMBRES[i],25])
    wb.save(path)


def crear_word(path, nombres=NOMBRES):
    doc=docx.Document();t=doc.add_table(rows=1,cols=2)
    t.rows[0].cells[0].text='Puesto';t.rows[0].cells[1].text='Nivel'
    for nombre in nombres:
        row=t.add_row();row.cells[0].text=nombre;row.cells[1].text='25'
    doc.save(path)


def test_ramas_completas_end_to_end_y_reanalisis(tmp_path):
    excel,word,org=(tmp_path/n for n in ['a.xlsx','a.docx','a.pdf'])
    crear_excel(excel,repetir_header=True);crear_word(word);crear_pdf(org)
    resultado=ejecutar_analisis(str(excel),str(org),str(word))
    assert resultado['errores_fatales']==[]
    assert [r.puesto_original for r in resultado['organigrama_records']]==NOMBRES
    assert all(not r.error_orden for r in resultado['organigrama_records'])
    assert all(r.estado_orden=='Misma posición' for r in resultado['resultados'])
    assert all(r.tipo_inconsistencia==TipoInconsistencia.OK for r in resultado['resultados'])
    cambiados=NOMBRES.copy();cambiados[2],cambiados[3]=cambiados[3],cambiados[2]
    crear_word(word,cambiados)
    movidos=ejecutar_analisis(str(excel),str(org),str(word))['resultados']
    assert [r.posiciones['Excel'] for r in movidos if r.estado_orden=='Diferente posición']==[[3],[4]]
    assert 'Word: posición 4; esperada 3' in movidos[2].observacion_orden
    crear_word(word)
    assert all(r.estado_orden=='Misma posición' for r in ejecutar_analisis(str(excel),str(org),str(word))['resultados'])


def test_id_11_se_conserva_aunque_sea_primera_fila(tmp_path):
    path=tmp_path/'id.xlsx';crear_excel(path,ids=list(range(11,19)))
    excel=extraer_excel(str(path))
    org=[PuestoRecord(Fuente.ORGANIGRAMA,r.puesto_original,'25',r.puesto_normalizado,'N25') for r in excel]
    resultados=comparar_fuentes(excel,[],org,incluir_word=False)
    assert resultados[0].posiciones['Excel']==[11]
    assert resultados[0].posiciones['Organigrama']==[1]
    assert resultados[0].estado_orden=='Diferente posición'
    assert 'esperada 11' in resultados[0].observacion_orden


def test_excel_id_ordenado_sin_usar_fila_fisica(tmp_path):
    path=tmp_path/'id.xlsx';crear_excel(path,ids=[8,7,6,5,4,3,2,1])
    registros=extraer_excel(str(path))
    assert [r.consecutivo_excel for r in registros]==list(range(1,9))
    assert registros[0].puesto_original==NOMBRES[-1]


@pytest.mark.parametrize('valor',[None,'PUESTO','abc',0,-1,1.5,True,'NaN','Infinity'])
def test_consecutivo_invalido_no_se_inventa(valor):
    assert _leer_consecutivo(valor) is None


@pytest.mark.parametrize('valor',[11,'11','011',11.0,'11.0'])
def test_consecutivo_entero(valor):
    assert _leer_consecutivo(valor)==11


def test_ids_repetidos_y_vacios_son_revision(tmp_path):
    path=tmp_path/'ids.xlsx';crear_excel(path,ids=[1,1,None,4,5,6,7,8])
    records=extraer_excel(str(path));bad=[r for r in records if r.error_orden]
    assert len(bad)==3 and all(r.valido for r in bad)
    resultados=comparar_fuentes(records,[],[],incluir_word=False)
    assert all(r.estado_orden=='No verificable' for r in resultados if r.excel.error_orden)
    assert all('Revisar consecutivo' in r.observacion_orden for r in resultados if r.excel.error_orden)


@pytest.mark.parametrize('opciones', [dict(conectado=False),dict(curva=True),dict(multipagina=True)])
def test_no_certifica_geometria_incompleta(tmp_path,opciones):
    path=tmp_path/'org.pdf';crear_pdf(path,**opciones)
    registros=extraer_organigrama(str(path))
    assert registros and all(r.error_orden for r in registros if r.valido)
    excel=[PuestoRecord(Fuente.EXCEL,r.puesto_original,'25',r.puesto_normalizado,'N25',consecutivo_excel=i) for i,r in enumerate(registros,1)]
    resultado=comparar_fuentes(excel,[],registros,incluir_word=False)
    assert all(r.estado_orden=='No verificable' for r in resultado)
    assert all(not r.posiciones['Organigrama'] for r in resultado)
    wb=openpyxl.load_workbook(BytesIO(generar_reporte_bytes(resultado,modo='excel_organigrama')))
    assert all(row[-3]=='No verificable' for row in wb['Detalle'].iter_rows(min_row=2,values_only=True))

def test_nombre_duplicado_con_un_id_vacio_no_se_certifica(tmp_path):
    path=tmp_path/'duplicado.xlsx';crear_excel(path)
    wb=openpyxl.load_workbook(path);ws=wb.active
    ws.append(['']*8+[None,NOMBRES[0],25]);wb.save(path)
    excel=extraer_excel(str(path))
    org=[PuestoRecord(Fuente.ORGANIGRAMA,r.puesto_original,'25',r.puesto_normalizado,'N25') for r in excel[:8]]
    principal=next(r for r in comparar_fuentes(excel,[],org,incluir_word=False)
                   if r.excel and r.excel.puesto_original==NOMBRES[0] and not r.puesto_clave_normalizada.startswith('[DUP-'))
    assert principal.estado_orden=='No verificable'


def test_conector_lateral_no_infiere_superior(tmp_path):
    path=tmp_path/'lateral.pdf';crear_pdf(path)
    doc=fitz.open(path);page=doc[0]
    page.draw_line((240,160),(470,160))
    out=tmp_path/'lateral2.pdf';doc.save(out);doc.close()
    assert all(r.error_orden for r in extraer_organigrama(str(out)))


def test_cruce_interior_de_conectores_no_es_union(tmp_path):
    path=tmp_path/'cruce.pdf';crear_pdf(path)
    doc=fitz.open(path);page=doc[0]
    page.draw_line((350,80),(350,120))  # atraviesa la barra horizontal sin terminar en ella
    out=tmp_path/'cruce2.pdf';doc.save(out);doc.close()
    assert all('Cruce' in r.error_orden for r in extraer_organigrama(str(out)))
